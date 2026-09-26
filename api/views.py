from django.contrib.auth.models import User

from django.shortcuts import render
from rest_framework import generics

from rest_framework.views import APIView

from collections import defaultdict

from .serializers import EventDetailSerializer
from .serializers import EventListSerializer
from .serializers import EventCardSerializer
from .serializers import ParticipantRegisteredEventSerializer
from .models import EventCategoryTable, EventDetailTable

from .serializers import EventSubDetailSerializer
from .models import EventSubDetailTable

from .serializers import ParticipantSerializer
from .serializers import ParticipantDetailSerializer
from .models import ParticipantTable

from .serializers import ParticipantPaymTableSerializer
from .models import ParticipantPaymRefTable

from .serializers import EventImageSerializer
from .models import EventImages

from .serializers import EventNotificationSerializer
from .serializers import OrganisationNotificationSerializer
from .models import EventNotificationTable
from .models import OrganisationNotificationTable

from .serializers import EventSponsorSerializer
from .models import EventSponsorTable

from .serializers import AppSponsorSerializer
from .models import AppSponsorTable


from .serializers import EventInformationSerializer
from .models import EventInformationTable

from .serializers import OrganisationSerializer
from .serializers import OrganisationMemberSerializer
from .serializers import OrganisationEventSerializer
from .serializers import AthleticOrganisationSerializer
from .serializers import AthleticOrganisationMemberSerializer
from .serializers import AthleticOrganisationCategorySerializer
from .serializers import ErrorSerializer
from .serializers import AppReleaseVersionSerializer
from .serializers import TermsAndConditionsSerializer
from .serializers import ParticipantTermsAcceptanceSerializer
from .serializers import ParticipantTermsStatusSerializer
from .models import OrganisationTable
from .models import OrganisationMemberTable
from .models import OrganisationEventTable
from .models import AthleticOrganisationTable
from .models import AthleticOrganisationMemberTable
from .models import AthleticOrganisationCategoryTable
from .models import ErrorTable
from .models import AppReleaseVersionTable
from .models import TermsAndConditionsTable
from .models import ParticipantTermsAcceptanceTable

from .models import ParticipantEventTable

from .util import EventStatus
from .util import OrganisationMemberStatus
from .util import OrganisationType
from django.db.models import Q
from django.db.models import Prefetch

import os
from decouple import config

from django.db import transaction

from django.utils import dateparse, timezone

#from django.http import JsonResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.decorators import api_view, permission_classes, action
from rest_framework.response import Response
from rest_framework import status
from rest_framework.viewsets import ModelViewSet 

from eventrunnerbe.notifications.utils import send_push_notification

from openpyxl import load_workbook

from api.storage_utils import generate_s3_presigned_url

#from eventrunnerbe.utils import customTokenBackend

#from eventrunnerbe.utils import CustomJWTAuthentication



from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
#from multipart import MultipartParser
#from multipart import Form

from functools import wraps

import jwt

import json
import logging

logger = logging.getLogger(__name__)


def _parse_request_body(request):
    try:
        if request.body:
            return json.loads(request.body.decode('utf-8'))
    except (json.JSONDecodeError, UnicodeDecodeError, AttributeError):
        logger.debug("Could not parse request body as JSON", extra={"path": request.path})
    return request.data if isinstance(request.data, dict) else {}


def _parse_body_or_400(request):
    if request.body:
        try:
            decoded = request.body.decode('utf-8')
            if decoded.strip():
                parsed = json.loads(decoded)
                if isinstance(parsed, dict):
                    return parsed, None
                logger.warning(
                    "JSON body is not an object",
                    extra={"path": request.path},
                )
                return None, Response(
                    {'detail': 'JSON body must be an object.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        except (json.JSONDecodeError, UnicodeDecodeError):
            logger.warning("Invalid JSON body", extra={"path": request.path})
            return None, Response(
                {'detail': 'Invalid JSON.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
    if isinstance(request.data, dict):
        return request.data, None
    return {}, None


def _queryset_by_event_or_id(queryset, body, event_field='event', id_field='id', ordering=None):
    if ordering is not None:
        queryset = queryset.order_by(ordering)
    if event_field:
        event_id = body.get(event_field)
        if event_id is not None and event_id != '':
            return queryset.filter(**{event_field: event_id}), True
    record_id = body.get(id_field)
    if record_id is not None and record_id != '':
        return queryset.filter(**{id_field: record_id}), False
    return queryset, True


def _filtered_viewset_response(
    request,
    queryset,
    serializer_class,
    event_field='event',
    id_field='id',
    ordering=None,
):
    body, error_response = _parse_body_or_400(request)
    if error_response:
        return error_response
    records, many = _queryset_by_event_or_id(
        queryset,
        body,
        event_field=event_field,
        id_field=id_field,
        ordering=ordering,
    )
    serializer = serializer_class(records, many=many)
    return Response(serializer.data, status=status.HTTP_200_OK)


def _get_participant_queryset_or_404(queryset, **lookup):
    try:
        return queryset.get(**lookup), None
    except ParticipantTable.DoesNotExist:
        logger.debug("Participant not found", extra=lookup)
        return None, Response(
            {'detail': 'Participant not found.'},
            status=status.HTTP_404_NOT_FOUND,
        )


def _get_event_or_404(event_id):
    try:
        return EventDetailTable.objects.get(id=event_id), None
    except EventDetailTable.DoesNotExist:
        logger.warning("Event not found", extra={"event_id": event_id})
        return None, Response(
            {'detail': 'Event not found.'},
            status=status.HTTP_404_NOT_FOUND,
        )


def _get_participant_or_404(participant_id):
    try:
        return ParticipantTable.objects.get(id=participant_id), None
    except ParticipantTable.DoesNotExist:
        logger.warning("Participant not found", extra={"participant_id": participant_id})
        return None, Response(
            {'detail': 'Participant not found.'},
            status=status.HTTP_404_NOT_FOUND,
        )


def _filter_queryset(queryset, body, filters):
    for key, field in filters.items():
        if key in body:
            return queryset.filter(**{field: body[key]})
    return queryset


def _get_participant_from_body(request):
    body = _parse_request_body(request)
    participant_id = body.get('participantid')
    if not participant_id:
        return None, Response(
            {'detail': 'participantid is required in request body.'},
            status=status.HTTP_400_BAD_REQUEST,
        )
    try:
        return ParticipantTable.objects.get(id=participant_id), None
    except ParticipantTable.DoesNotExist:
        logger.debug("Participant not found from request body", extra={"participant_id": participant_id})
        return None, Response(
            {'detail': 'Participant not found.'},
            status=status.HTTP_404_NOT_FOUND,
        )


def _completed_event_ids_for_participant(participant):
    return ParticipantEventTable.objects.filter(
        participant=participant,
        paymref__payment_status='Completed',
    ).values_list('event_id', flat=True).distinct()


def _event_card_queryset():
    return EventDetailTable.objects.select_related(
        'athleticscategory',
        'athleticorganisation',
        'eventimage_event',
    )


def _get_current_terms(document_type='app'):
    return TermsAndConditionsTable.objects.filter(
        document_type=document_type,
        is_current=True,
    ).first()


def _build_participant_terms_status(participant, document_type='app'):
    current_terms = _get_current_terms(document_type)
    current_acceptance = None
    if current_terms:
        current_acceptance = ParticipantTermsAcceptanceTable.objects.filter(
            participant=participant,
            terms=current_terms,
        ).first()

    last_acceptance = ParticipantTermsAcceptanceTable.objects.filter(
        participant=participant,
        terms__document_type=document_type,
    ).select_related('terms').order_by('-accepted_at').first()

    return {
        'accepted': current_acceptance is not None,
        'document_type': document_type,
        'current_terms_id': current_terms.id if current_terms else None,
        'current_version': current_terms.version_number if current_terms else None,
        'current_terms': current_terms if current_acceptance is not None else None,
        'last_accepted_terms_id': last_acceptance.terms_id if last_acceptance else None,
        'last_accepted_version': last_acceptance.terms.version_number if last_acceptance else None,
        'last_accepted_at': last_acceptance.accepted_at if last_acceptance else None,
    }


def _participant_detail_queryset():
    return ParticipantTable.objects.prefetch_related(
        Prefetch(
            'athleticorganisationmember_participant',
            queryset=AthleticOrganisationMemberTable.objects.filter(
                status=OrganisationMemberStatus.Active,
            ).select_related('athleticorganisation'),
        ),
    )


ATHLETIC_ORGANISATION_MOVE_POSITIONS = {
    'first-child', 'last-child', 'left', 'right', 'sorted-child',
}


def _coerce_bool(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return None
    return str(value).lower() in ('1', 'true', 'yes', 'on')


def _athletic_organisation_kwargs(data):
    scalar_fields = [
        'code', 'name', 'registrationnumber', 'countrycode', 'province', 'postcode',
        'city', 'physicaladdress', 'contactname', 'contactsurname', 'contactemail',
        'contactphone', 'websiteurl', 'lat', 'lng',
    ]
    kwargs = {}
    for field in scalar_fields:
        if field in data and data[field] not in (None, ''):
            kwargs[field] = data[field]

    if 'organisationtype' in data and data['organisationtype'] not in (None, ''):
        kwargs['organisationtype'] = int(data['organisationtype'])

    for bool_field in ('isverified', 'isactive'):
        if bool_field in data:
            coerced = _coerce_bool(data[bool_field])
            if coerced is not None:
                kwargs[bool_field] = coerced

    return kwargs


def _validate_athletic_organisation_required_fields(kwargs):
    missing = []
    if not kwargs.get('code'):
        missing.append('code')
    if not kwargs.get('name'):
        missing.append('name')
    if kwargs.get('organisationtype') is None:
        missing.append('organisationtype')
    return missing


def _move_athletic_organisation(organisation, target_parent=None, position='sorted-child'):
    if position not in ATHLETIC_ORGANISATION_MOVE_POSITIONS:
        raise ValueError(f'Invalid position. Use one of: {", ".join(sorted(ATHLETIC_ORGANISATION_MOVE_POSITIONS))}')

    if target_parent is not None:
        if organisation.pk == target_parent.pk:
            raise ValueError('An organisation cannot be moved under itself.')
        if target_parent.is_descendant_of(organisation):
            raise ValueError('An organisation cannot be moved under its own descendant.')
        organisation.move(target_parent, position)
        return organisation

    if organisation.is_root():
        return organisation

    first_root = AthleticOrganisationTable.get_first_root_node()
    if first_root is None or first_root.pk == organisation.pk:
        return organisation

    organisation.move(first_root, 'left')
    return organisation


def _update_athletic_organisation_fields(organisation, data):
    kwargs = _athletic_organisation_kwargs(data)
    kwargs.pop('code', None)
    for field, value in kwargs.items():
        setattr(organisation, field, value)
    organisation.save()
    return organisation

import requests

from django.http import JsonResponse


def get_token_auth_header(request):
    """Obtains the Access Token from the Authorization Header"""
    auth = request.META.get("HTTP_AUTHORIZATION")
    if not auth:
        logger.warning("Missing Authorization header")
        raise ValueError("Authorization header is missing.")

    parts = auth.split()
    if len(parts) != 2 or parts[0].lower() != "bearer":
        logger.warning("Malformed Authorization header")
        raise ValueError("Authorization header must be Bearer token.")

    return parts[1]


def jwt_decode_token(token):
    try:
        header = jwt.get_unverified_header(token)
    except jwt.PyJWTError:
        logger.warning("Invalid JWT token header")
        raise ValueError("Invalid token.") from None

    auth0_domain = 'dev-x8hbr3jrn2mxvw4x.us.auth0.com'
    try:
        jwks_response = requests.get(
            f'https://{auth0_domain}/.well-known/jwks.json',
            timeout=10,
        )
        jwks_response.raise_for_status()
        jwks = jwks_response.json()
    except requests.RequestException:
        logger.exception("Failed to fetch Auth0 JWKS")
        raise ValueError("Unable to validate token.") from None

    public_key = None
    for jwk in jwks.get('keys', []):
        if jwk.get('kid') == header.get('kid'):
            public_key = jwt.algorithms.RSAAlgorithm.from_jwk(json.dumps(jwk))
            break

    if public_key is None:
        logger.error("Auth0 public key not found", extra={"kid": header.get('kid')})
        raise ValueError('Public key not found.')

    issuer = f'https://{auth0_domain}/'
    try:
        return jwt.decode(
            token,
            public_key,
            audience='https://eventrunner.com/api/',
            issuer=issuer,
            algorithms=['RS256'],
        )
    except jwt.PyJWTError:
        logger.warning("JWT decode failed")
        raise ValueError("Invalid token.") from None


def requires_scope(required_scope):
    """Determines if the required scope is present in the Access Token"""
    def require_scope(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            try:
                token = get_token_auth_header(args[0])
                decoded = jwt_decode_token(token)
            except ValueError as exc:
                response = JsonResponse({'message': str(exc)})
                response.status_code = status.HTTP_401_UNAUTHORIZED
                return response

            if decoded.get("scope"):
                token_scopes = decoded["scope"].split()
                for token_scope in token_scopes:
                    if token_scope == required_scope:
                        return f(*args, **kwargs)

            logger.warning(
                "Insufficient token scope",
                extra={"required_scope": required_scope},
            )
            response = JsonResponse({'message': 'You don\'t have access to this resource'})
            response.status_code = 403
            return response
        return decorated
    return require_scope


#class EventDetailListCreate(generics.ListCreateAPIView)
#    queryset = EventDetailTable.objects.all()
#    serializer_class = eventdetailserializer

@api_view(['GET'])
def getRoutes(request):
    routes = [
        {
            'Endpoint' : '/events/',
            'method' : 'GET',
            'body' : None,            
            'description' : 'Returns an array of events'
        },
        {
            'Endpoint' : '/events/id',
            'method' : 'GET',
            'body' : None,            
            'description' : 'Returns a single event'
        }
    ]
    return Response(routes)

##### Event details
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def getEvents(request):
    eventdetails = EventDetailTable.objects.all().order_by('-eventdate')
    eventserializer = EventDetailSerializer(eventdetails , many=True)
    return  Response(eventserializer.data)

@api_view(['GET'])
def getEventsList(request):
    eventdetails = EventDetailTable.objects.only(
        'id', 'eventname', 'eventdate', 'eventenddate', 'eventstatus',
        'eventtype', 'contactpersonname', 'contactpersonemailaddr',
    ).order_by('-eventdate')
    eventserializer = EventListSerializer(eventdetails, many=True)
    return Response(eventserializer.data)


@api_view(['POST'])
def getParticipantRegisteredUpcomingEventCards(request):
    participant, error_response = _get_participant_from_body(request)
    if error_response:
        return error_response

    completed_event_ids = _completed_event_ids_for_participant(participant)
    events = _event_card_queryset().filter(
        isactive=True,
        eventstatus__in=[EventStatus.RegistrationOpen, EventStatus.RegistrationClosed],
        id__in=completed_event_ids,
    ).order_by('eventdate')

    return Response(EventCardSerializer(events, many=True).data)


@api_view(['POST'])
def getParticipantAvailableUpcomingEventCards(request):
    participant, error_response = _get_participant_from_body(request)
    if error_response:
        return error_response

    completed_event_ids = _completed_event_ids_for_participant(participant)
    events = _event_card_queryset().filter(
        isactive=True,
        eventstatus__in=[EventStatus.RegistrationOpen, EventStatus.RegistrationClosed],
    ).exclude(id__in=completed_event_ids).order_by('eventdate')

    return Response(EventCardSerializer(events, many=True).data)


@api_view(['POST'])
def getParticipantCompletedPastEventCards(request):
    participant, error_response = _get_participant_from_body(request)
    if error_response:
        return error_response

    completed_event_ids = _completed_event_ids_for_participant(participant)
    events = _event_card_queryset().filter(
        isactive=True,
        eventstatus__in=[EventStatus.Closed, EventStatus.Completed],
        id__in=completed_event_ids,
    ).order_by('-eventdate')

    return Response(EventCardSerializer(events, many=True).data)


@api_view(['GET'])
#@requires_scope('read:events')
def getEvent(request , pk):
    eventdetail, error_response = _get_event_or_404(pk)
    if error_response:
        return error_response
    eventserializer = EventDetailSerializer(eventdetail , many=False)
    return Response(eventserializer.data)

@api_view(['GET'])
#@requires_scope('read:events')
def getEventRelated(request , pk):
    _, error_response = _get_event_or_404(pk)
    if error_response:
        return error_response

    eventImg = EventImages.objects.select_related('event').all()
    eventImgserializer = EventImageSerializer(eventImg , many=True)
    return Response(eventImgserializer.data)

@api_view(['POST'])
#@requires_scope('create:event')
def createEvent(request):
    data = request.data
    try:
        eventdetail = EventDetailTable.objects.create(
            eventid=data['eventid'],
            eventname=data['eventname'],
            contactpersonname=data['contactpersonname'],
            contactpersonsurname=data['contactpersonsurname'],
            eventstatus=data['eventstatus'],
            eventdescription=data['eventdescription'],
            eventtype=data['eventtype'],
            eventcategory=data['eventcategory'],
        )
    except KeyError as exc:
        logger.warning("createEvent missing required field", extra={"field": str(exc)})
        return Response(
            {'detail': f'Missing required field: {exc}'},
            status=status.HTTP_400_BAD_REQUEST,
        )
    except Exception:
        logger.exception("createEvent failed")
        return Response(
            {'detail': 'Could not create event.'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    logger.info("Event created", extra={"event_id": eventdetail.id})
    eventserializer = EventDetailSerializer(eventdetail , many=False)
    return Response(eventserializer.data)

@api_view(['PUT'])
#@requires_scope('update:event')
def updateEvent(request , pk):
    eventdetail, error_response = _get_event_or_404(pk)
    if error_response:
        return error_response

    eventserializer = EventDetailSerializer(eventdetail , data=request.data)
    if eventserializer.is_valid():
        eventserializer.save()
        logger.info("Event updated", extra={"event_id": pk})
    else:
        logger.warning(
            "updateEvent validation failed",
            extra={"event_id": pk, "errors": eventserializer.errors},
        )
    return Response(eventserializer.data)

@api_view(['DELETE'])
#@requires_scope('delete:event')
def deleteEvent(request , pk):
    eventdetail, error_response = _get_event_or_404(pk)
    if error_response:
        return error_response

    eventdetail.delete()
    logger.info("Event deleted", extra={"event_id": pk})
    return Response("Event was deleted")

@api_view(['GET'])
#@requires_scope('read:events')
def getEventsUnregistered(request , pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response

    eventlist = ParticipantEventTable.objects.filter(
        participant=participantdetail,
        paymref__payment_status='Completed',
    ).values('event')
    eventdetail = EventDetailTable.objects.exclude(id__in=eventlist).exclude(
        Q(eventstatus=EventStatus.Completed) | Q(eventstatus=EventStatus.Closed),
    ).order_by('eventdate')
    eventserializer = EventDetailSerializer(eventdetail , many=True)
    return Response(eventserializer.data)

##### Participants
@api_view(['GET'])
#@requires_scope('read:profile')
def getParticipants(request):
    participantdetails = ParticipantTable.objects.all()
    participantserializer = ParticipantSerializer(participantdetails , many=True)
    return  Response(participantserializer.data)

@api_view(['GET'])
#@requires_scope('read:profile')
def getParticipant(request , pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response
    participantserializer = ParticipantSerializer(participantdetail , many=False)
    return Response(participantserializer.data)

@api_view(['GET'])
#@requires_scope('read:profile')
def getParticipantExists(request , pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response
    participantserializer = ParticipantSerializer(participantdetail , many=False)

    if participantdetail.authid:
        return Response("Participant exists" , status=status.HTTP_200_OK)
    return Response(status=status.HTTP_204_NO_CONTENT)

@api_view(['POST'])
#@requires_scope('create:profile')
def createParticipant(request):
    data = request.data
    logger.debug("createParticipant request received", extra={"field_count": len(data)})
    participantdetail = ParticipantTable.objects.create(
        user = data['user'],
        title = data['title'],
        surname = data['surname'],
        firstname = data['firstname'],
        othernames = data['othernames'],
        initials = data['initials'],
        preferredname = data['preferredname'],
        homelanguage = data['homelanguage'],
        preferredlanguage = data['preferredlanguage'],
        maidenname = data['maidenname'],
        countryofissue = data['countryofissue'],
        typefld = data['typefld'],
        disabled = data['disabled']
    )
    participantserializer = ParticipantSerializer(participantdetail , many=False)
    return  Response(participantserializer.data)

@api_view(['PUT'])
#@requires_scope('update:profile')
def updateParticipant(request , pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response

    participantserializer = ParticipantSerializer(participantdetail , data=request.data)
    if participantserializer.is_valid():
        participantserializer.save()
        logger.info("Participant updated", extra={"participant_id": pk})
    else:
        logger.warning(
            "updateParticipant validation failed",
            extra={"participant_id": pk, "errors": participantserializer.errors},
        )
    return Response(participantserializer.data)

@api_view(['DELETE'])
#@requires_scope('delete:profile')
def deleteParticipant(request , pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response

    participantdetail.delete()
    logger.info("Participant deleted", extra={"participant_id": pk})
    return Response("Participant was deleted")

##### Participants  events
@api_view(['GET'])
#@requires_scope('read:events')
def getParticipantEvents(request, pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response

    participant_events = ParticipantEventTable.objects.filter(
        participant=participantdetail,
        paymref__payment_status='Completed',
    ).values_list('event_id', 'subevent_id').distinct()

    subevents_by_event = defaultdict(list)
    event_ids = set()
    for event_id, subevent_id in participant_events:
        event_ids.add(event_id)
        subevents_by_event[event_id].append(subevent_id)

    #all_subevent_ids = [sid for ids in subevents_by_event.values() for sid in ids]
    #subevents_map = {
    #    subevent.id: subevent
    #    for subevent in EventSubDetailTable.objects.filter(id__in=all_subevent_ids)
    #}


    eventRegDetails = EventDetailTable.objects.filter(
        id__in=event_ids,
    ).filter(eventstatus__in=[EventStatus.RegistrationOpen, EventStatus.RegistrationClosed]
    ).prefetch_related(
        'subevent_event',
        'eventimage_event',
        'eventnotification_event',
        'eventsponsor_event',
    ).select_related(
        'organisation',
        'athleticscategory',
    ).order_by('eventdate')

    eventserializer = ParticipantRegisteredEventSerializer(
        eventRegDetails,
        many=True,
        context={
            'subevents_by_event': dict(subevents_by_event),
        #    'subevents_map': subevents_map,
        },
    )
    return Response(eventserializer.data)

@api_view(['GET'])
#@requires_scope('read:events')
def getEventsParticipant(request, pk):
    eventDetTable, error_response = _get_event_or_404(pk)
    if error_response:
        return error_response
    
    participantslist = ParticipantEventTable.objects.filter(
        event=eventDetTable,
        paymref__payment_status='Completed',
    ).values('participant')
    particpantsDetails = ParticipantTable.objects.filter(id__in=participantslist)
    particpantsserializer = ParticipantSerializer(particpantsDetails , many=True)
    return Response(particpantsserializer.data)

@api_view(['GET'])
#@requires_scope('read:events')
def getParticipantCompletedEvents(request, pk):
    participantdetail, error_response = _get_participant_or_404(pk)
    if error_response:
        return error_response
    #eventdetails = participantdetail.events.filter(eventstatus__in=[EventStatus.Closed, EventStatus.Completed])   
    #eventserializer = EventDetailSerializer(eventdetails , many=True)

    eventlist = ParticipantEventTable.objects.filter(participant=participantdetail).filter(paymref__payment_status='Completed').values('event')

    eventRegDetails = EventDetailTable.objects.filter(id__in=eventlist).filter(eventstatus__in=[EventStatus.Closed, EventStatus.Completed]).order_by('-eventdate')
    eventserializer = EventDetailSerializer(eventRegDetails , many=True)

    return Response(eventserializer.data)

@api_view(['POST'])
#@requires_scope('update:participant')
def updateParticipantEvent(request):
    data = request.data
    participant_id = data.get('user')
    event_id = data.get('event')
    if not participant_id or not event_id:
        return Response(
            {'detail': 'user and event are required.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    participantdetail, error_response = _get_participant_or_404(participant_id)
    if error_response:
        return error_response

    eventdetail, error_response = _get_event_or_404(event_id)
    if error_response:
        return error_response

    participantdetail.events.add(eventdetail)
    logger.info(
        "Participant enrolled in event",
        extra={"participant_id": participant_id, "event_id": event_id},
    )
    return Response("Participant was enrolled to the event")

@api_view(['POST'])
#@requires_scope('update:event')
def removeParticipantEvent(request):
    data = request.data
    participant_id = data.get('user')
    event_id = data.get('event')
    if not participant_id or not event_id:
        return Response(
            {'detail': 'user and event are required.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    participantdetail, error_response = _get_participant_or_404(participant_id)
    if error_response:
        return error_response

    eventdetail, error_response = _get_event_or_404(event_id)
    if error_response:
        return error_response

    participantdetail.events.remove(eventdetail)
    logger.info(
        "Participant removed from event",
        extra={"participant_id": participant_id, "event_id": event_id},
    )
    return Response("Participant was disenrolled to the event")

##### Events images
@api_view(['GET'])
def getEventImage(request , pk):
    #eventDet = EventDetailTable.objects.get(id = pk)
    eventimages = EventImages.objects.filter(event=pk)
    if eventimages.count():
        respResult = "Count is not zero"
    else:
        respResult = "Count is zero"

    eventimgserializer = EventImageSerializer(eventimages , many=True)
    #return  Response(eventimgserializer.data)
    return  Response(eventimgserializer.data)

##### Events images
@api_view(['GET'])
def getEventImageUsable(request , pk):
    eventimages = EventImages.objects.filter(event=pk)
    data = []

    for obj in eventimages:
        presigned_url = generate_s3_presigned_url('media/public/' + obj.eventMainImg.name)

        data.append({
            'id': obj.id,
            'event': obj.event.id,
            'eventMainImg': presigned_url if presigned_url else None,
        })

    return Response(data, status=status.HTTP_200_OK)

            
class eventImageViewSet(ModelViewSet):
    queryset = EventImages.objects.all()
    serializer_class = EventImageSerializer
    parser_classes = (MultiPartParser , FormParser)

    def create(self , request  , *args, **kwargs):
        event = request.data.get('event')
        if not event:
            return Response({'detail': 'event is required.'}, status=status.HTTP_400_BAD_REQUEST)

        eventDetTable, error_response = _get_event_or_404(event)
        if error_response:
            return error_response

        eventMainImg = request.data.get('eventMainImg')
        if not eventMainImg:
            return Response({'detail': 'eventMainImg is required.'}, status=status.HTTP_400_BAD_REQUEST)

        eventImg = EventImages.objects.filter(event=event)
        if eventImg.count() != 0:
            eventImg.delete()

        EventImages.objects.create(event=eventDetTable, eventMainImg=eventMainImg)
        logger.info("Event image uploaded", extra={"event_id": event})

        return Response("Event image uploaded successfully", status=status.HTTP_201_CREATED)
    
    def retrieve(self , request  , *args, **kwargs):
        event = request.data.get('event')
        if not event:
            return Response({'detail': 'event is required.'}, status=status.HTTP_400_BAD_REQUEST)

        _, error_response = _get_event_or_404(event)
        if error_response:
            return error_response

        eventImg = EventImages.objects.filter(event=event)
        eventimgserializer = EventImageSerializer(eventImg , many=True)
        return Response(eventimgserializer.data)

    
class participantCheckViewSet(ModelViewSet):
    queryset = ParticipantTable.objects.all()
    serializer_class = ParticipantSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        participantAuthid = request.data.get('authid')
        reqType = request.data.get('requesttype')

        if not participantAuthid:
            return Response({'detail': 'authid is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            participant = ParticipantTable.objects.get(authid=participantAuthid)
        except ParticipantTable.DoesNotExist:
            logger.debug("Participant check not found", extra={"authid": participantAuthid})
            return Response(status=status.HTTP_204_NO_CONTENT)

        if reqType == 'check':
            if participant.authid:
                return Response("Participant exists", status=status.HTTP_200_OK)
            return Response(status=status.HTTP_204_NO_CONTENT)

        participantserializer = ParticipantSerializer(participant , many=False)
        return Response(participantserializer.data, status=status.HTTP_200_OK)
        
    def list(self , request  , *args, **kwargs):
        participantAuthid = request.data.get('authid')
        reqType = request.data.get('requesttype')

        if not participantAuthid:
            return Response({'detail': 'authid is required.'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            participant = ParticipantTable.objects.get(authid = participantAuthid)
        except ParticipantTable.DoesNotExist:
            logger.debug("Participant check not found", extra={"authid": participantAuthid})
            return Response("")

        if  reqType == 'check':
            return  Response(participant.id)
        else:
            participantserializer = ParticipantSerializer(participant , many=False)
            return  Response(participantserializer.data)
        
class participantViewSet(ModelViewSet):
    queryset = _participant_detail_queryset()
    serializer_class = ParticipantSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def get_serializer_class(self):
        if self.action in ('list', 'retrieve', 'update', 'partial_update', 'lookup_by_email'):
            return ParticipantDetailSerializer
        return ParticipantSerializer

    def retrieve(self , request  , *args, **kwargs):
        participantEmailAddr = request.data.get('emailaddress')
        reqType = request.data.get('reqtype')

        if not participantEmailAddr:
            return Response({'detail': 'emailaddress is required.'}, status=status.HTTP_400_BAD_REQUEST)

        participant, error_response = _get_participant_queryset_or_404(
            _participant_detail_queryset(),
            emailaddress=participantEmailAddr,
        )
        if error_response:
            if reqType == 'check':
                return Response(status=status.HTTP_204_NO_CONTENT)
            return error_response

        if reqType == 'check':
            return Response("Participant exists", status=status.HTTP_200_OK)

        participantserializer = ParticipantDetailSerializer(participant , many=False)
        return Response(participantserializer.data, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post'], url_path='lookup-by-email')
    def lookup_by_email(self, request):
        body = _parse_request_body(request)
        participant_email = body.get('emailaddress')
        if not participant_email:
            return Response(
                {'detail': 'emailaddress is required in request body.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            participant = _participant_detail_queryset().get(emailaddress=participant_email)
        except ParticipantTable.DoesNotExist:
            logger.debug("Participant lookup by email not found")
            return Response(status=status.HTTP_204_NO_CONTENT)

        return Response(
            ParticipantDetailSerializer(participant).data,
            status=status.HTTP_200_OK,
        )
        
    def create(self , request  , *args, **kwargs):

        is_many = isinstance(request.data, list)
        if (is_many):
            serializer = self.get_serializer(data=request.data, many=is_many)
            serializer.is_valid(raise_exception=True)
            self.perform_create(serializer)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
    
        try:
            participanrExists = ParticipantTable.objects.get(emailaddress= request.data['emailaddress'])   
            return Response(participanrExists.id, status=status.HTTP_200_OK)
        except ParticipantTable.DoesNotExist:
            pass

        try:
            usrname = request.data["user"]
            usr = User.objects.get(username= usrname)             
        except User.DoesNotExist:
            logger.warning("Participant create failed: invalid user")
            return Response("Invalid user", status=status.HTTP_409_CONFLICT)
        
        if (request.data['emailaddress'] == None or request.data['emailaddress'] == ''):
            return Response("Email address is required", status=status.HTTP_400_BAD_REQUEST)
        
        participantTable = ParticipantTable.objects.create(
            user = usr,
            title = request.data['title'],
            surname = request.data['surname'],
            firstname = request.data['firstname'],
            othernames = request.data['othernames'],
            initials = request.data['initials'],
            preferredname = request.data['preferredname'],
            homelanguage = request.data['homelanguage'],
            preferredlanguage = request.data['preferredlanguage'],
            maidenname = request.data['maidenname'],
            countryofissue = request.data['countryofissue'],
            typefld = request.data['typefld'],
            emailaddress = request.data['emailaddress'],
            usrphonenum = request.data['usrphonenum'],
            authid = request.data['authid']
        )      

        reqData = request.data

        logger.debug(
            "participantViewSet.create completed",
            extra={"participant_id": participantTable.id},
        )

        participantserializer = ParticipantSerializer(participantTable , many=False)
        return  Response(participantserializer.data , status=status.HTTP_201_CREATED)
    
    def update(self , request  , *args, **kwargs):
        participantId = request.data.get('id')
        if not participantId:
            return Response({'detail': 'id is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            participantTable = ParticipantTable.objects.prefetch_related(
                Prefetch(
                    'athleticorganisationmember_participant',
                    queryset=AthleticOrganisationMemberTable.objects.filter(
                        status=OrganisationMemberStatus.Active,
                    ).select_related('athleticorganisation'),
                ),
            ).get(id=participantId)
        except ParticipantTable.DoesNotExist:
            logger.warning("Participant update failed: not found", extra={"participant_id": participantId})
            return Response({'detail': 'Participant not found.'}, status=status.HTTP_404_NOT_FOUND)

        reqData = request.data

        logger.debug(
            "participantViewSet.update",
            extra={"participant_id": participantId, "field_count": len(reqData)},
        )

        try:
            for reqAttr in request.data:
                if reqAttr == 'title':
                    participantTable.title = request.data[reqAttr]
                elif reqAttr == 'surname':
                    participantTable.surname = request.data[reqAttr]
                elif reqAttr == 'firstname':
                    participantTable.firstname = request.data[reqAttr]
                elif reqAttr == 'initials':
                    participantTable.initials = request.data[reqAttr]
                elif reqAttr == 'preferredname':
                    participantTable.preferredname = request.data[reqAttr]
                elif reqAttr == 'homelanguage':
                    participantTable.homelanguage = request.data[reqAttr]
                elif reqAttr == 'preferredlanguage':
                    participantTable.preferredlanguage = request.data[reqAttr]
                elif reqAttr == 'maidenname':
                    participantTable.maidenname = request.data[reqAttr]
                elif reqAttr == 'countryofissue':
                    participantTable.countryofissue = request.data[reqAttr]
                elif reqAttr == 'usrphonenum':
                    participantTable.usrphonenum = request.data[reqAttr]
                elif reqAttr == 'dateOfBirth':
                    datetime_object = dateparse.parse_datetime(request.data[reqAttr])
                    participantTable.dateOfBirth = datetime_object.date()
                elif reqAttr == 'gender':
                    try:
                        genderIntValue = int(request.data[reqAttr])
                    except (ValueError, TypeError) as exc:
                        raise ValueError("Invalid value for gender field. Must be an integer.") from exc
                    participantTable.gender = genderIntValue
                elif reqAttr == 'profilepic':
                    participantTable.profilepic = request.data[reqAttr]
                elif reqAttr == 'event':
                    event = request.data[reqAttr]
                    try:
                        eventDetTable = EventDetailTable.objects.get(id=event)
                        participantTable.events.add(eventDetTable)
                    except EventDetailTable.DoesNotExist:
                        return Response("Event does not exists", status=status.HTTP_400_BAD_REQUEST)
                elif reqAttr == 'user':
                    usr = User.objects.filter(username=request.data[reqAttr])
                    if usr.count() == 0:
                        return Response("User does not exists", status=status.HTTP_400_BAD_REQUEST)
                    participantTable.user = usr.first()
        except ValueError as exc:
            logger.warning(
                "Participant update validation failed",
                extra={"participant_id": participantId, "error": str(exc)},
            )
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        logger.debug(
            "participantViewSet.update saving",
            extra={"participant_id": participantId},
        )

        participantTable.save()
        participantserializer = ParticipantDetailSerializer(participantTable , many=False)
        return Response(participantserializer.data , status=status.HTTP_200_OK)   
   
    def partial_update(self , request  , *args, **kwargs):
        participantId = request.data.get('id')
        if not participantId:
            return Response({'detail': 'id is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            participantTable = ParticipantTable.objects.prefetch_related(
                Prefetch(
                    'athleticorganisationmember_participant',
                    queryset=AthleticOrganisationMemberTable.objects.filter(
                        status=OrganisationMemberStatus.Active,
                    ).select_related('athleticorganisation'),
                ),
            ).get(id=participantId)
        except ParticipantTable.DoesNotExist:
            logger.warning(
                "Participant partial update failed: not found",
                extra={"participant_id": participantId},
            )
            return Response({'detail': 'Participant not found.'}, status=status.HTTP_404_NOT_FOUND)

        participantTable.title = request.data.get('title', participantTable.title)
   
        participantserializer = ParticipantSerializer(participantTable , data=request.data)
        if participantserializer.is_valid():
            return Response("In valid")       
        return Response(participantId)

class eventViewSet(ModelViewSet):
    queryset = EventDetailTable.objects.all()
    serializer_class = EventDetailSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        eventid = request.data['id']        

        eventDetTable = EventDetailTable.objects.filter(id = eventid)    

        events = EventDetailTable.objects.prefetch_related("subevent").filter(id=eventid)   
        serializer = EventDetailSerializer(events, many=True)

        eventserializer = EventDetailSerializer(eventDetTable , many=True)
        return  Response(serializer.data, status=status.HTTP_200_OK)
    
    def list(self , request  , *args, **kwargs):
        try :
            eventid = request.data['id']
            events = EventDetailTable.objects.prefetch_related("eventimage_event").prefetch_related("subevent_event").prefetch_related("eventnotification_event").prefetch_related("eventsponsor_event").filter(id=eventid)  
                
        except KeyError:
            events = EventDetailTable.objects.prefetch_related("eventimage_event").prefetch_related("subevent_event").prefetch_related("eventnotification_event").prefetch_related("eventsponsor_event").all()
        except AttributeError:
            events = EventDetailTable.objects.prefetch_related("eventimage_event").prefetch_related("subevent_event").prefetch_related("eventnotification_event").prefetch_related("eventsponsor_event").all()
        #eventDetTable = EventDetailTable.objects.get(id = eventid)    

        serializer = EventDetailSerializer(events, many=True)

        #eventserializer = EventDetailSerializer(eventDetTable , many=False)
        return  Response(serializer.data, status=status.HTTP_200_OK)



class participantPaymViewSet(ModelViewSet):
    queryset = ParticipantPaymRefTable.objects.all()
    serializer_class = ParticipantPaymTableSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)


    def create(self , request  , *args, **kwargs):
        payment_context = {
            "participant_id": request.data.get("participant"),
            "event_id": request.data.get("event"),
            "subevent_id": request.data.get("subevent"),
            "payment_status": request.data.get("payment_status"),
        }

        try:
            participantTable = ParticipantTable.objects.get(id= request.data['participant'])   
        except ParticipantTable.DoesNotExist:
            logger.warning("Payment create failed: participant not found", extra=payment_context)
            return Response("Participant not found", status=status.HTTP_404_NOT_FOUND)
        
        try:
            eventTable = EventDetailTable.objects.get(id= request.data['event'])   
        except EventDetailTable.DoesNotExist:
            logger.warning("Payment create failed: event not found", extra=payment_context)
            return Response("Event not found", status=status.HTTP_404_NOT_FOUND)
        
        try:
            eventSubTable = EventSubDetailTable.objects.get(id= request.data['subevent'])   
        except EventSubDetailTable.DoesNotExist:
            logger.warning("Payment create failed: subevent not found", extra=payment_context)
            return Response("Sub event not found", status=status.HTTP_404_NOT_FOUND)

        try:
            with transaction.atomic():
                participantPaymRefTable = ParticipantPaymRefTable.objects.create(
                    participant = participantTable,
                    merchant_id = request.data['merchant_id'],
                    name_first = request.data['name_first'],
                    name_last = request.data['name_last'],
                    email_address = request.data['email_address'],
                    m_paym_id = request.data['m_payment_id'],
                    amount = request.data['amount'],
                    item_name = request.data['item_name'],
                    payment_uuid = request.data['payment_uuid'],
                    payment_timestamp = request.data['payment_timestamp'],
                    payment_status = request.data['payment_status']
                )    

                participant_event_id = None
                if request.data['payment_status'] == 'Completed':
                    participantEventTable = ParticipantEventTable.objects.create(
                        participant = participantTable,
                        event = eventTable,
                        subevent = eventSubTable,
                        paymref = participantPaymRefTable
                    )
                    participant_event_id = participantEventTable.id
            
                participantpaymserializer = ParticipantPaymTableSerializer(participantPaymRefTable , many=False)

        except Exception:
            logger.exception("Payment reference creation failed", extra=payment_context)
            return Response("Error creating payment reference", status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        logger.info(
            "Payment reference created",
            extra={
                **payment_context,
                "payment_ref_id": participantPaymRefTable.id,
                "participant_event_id": participant_event_id,
            },
        )

        return  Response(participantpaymserializer.data , status=status.HTTP_201_CREATED)

        
    
class EventSubViewSet(ModelViewSet):
    queryset = EventSubDetailTable.objects.all()
    serializer_class = EventSubDetailSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        eventid = request.data.get('event')
        if not eventid:
            return Response({'detail': 'event is required.'}, status=status.HTTP_400_BAD_REQUEST)

        eventsubDetTable = EventSubDetailTable.objects.filter(event=eventid)
        eventsubserializer = EventSubDetailSerializer(eventsubDetTable , many=True)
        return Response(eventsubserializer.data, status=status.HTTP_200_OK)
    
    def list(self , request  , *args, **kwargs):
        body, error_response = _parse_body_or_400(request)
        if error_response:
            return error_response

        eventid = body.get('event')
        if not eventid:
            return Response({'detail': 'event is required.'}, status=status.HTTP_400_BAD_REQUEST)

        eventsubDetTable = EventSubDetailTable.objects.filter(event=eventid)
        eventsubserializer = EventSubDetailSerializer(eventsubDetTable , many=True)
        return Response(eventsubserializer.data, status=status.HTTP_200_OK)
    
@api_view(['GET'])
def getSubeventsEvent(request , pk):
    subevents = EventSubDetailTable.objects.filter(event=pk)
    subeventserializer = EventSubDetailSerializer(subevents , many=True)

    return Response(subeventserializer.data, status=status.HTTP_200_OK)


@api_view(['GET'])
def getPayfastConnectionDetails(request):

    usesandbox = config('PAYFAST_USESANDBOX', cast=bool, default = True)
    prefix = 'PAYFAST_SANDBOX_' if usesandbox else 'PAYFAST_'
    env_keys = [
        f'{prefix}MERCHANT_ID',
        f'{prefix}MERCHANT_KEY',
        f'{prefix}MERCHANT_PASSPHRASE',
        f'{prefix}MERCHANT_ACTIVESCRIPT',
        f'{prefix}MERCHANT_PAYMMETHODS',
    ]
    missing_keys = [key for key in env_keys if not os.getenv(key)]
    if missing_keys:
        logger.warning(
            "PayFast configuration incomplete",
            extra={"usesandbox": usesandbox, "missing_keys": missing_keys},
        )

    merchantid = os.getenv('PAYFAST_SANDBOX_MERCHANT_ID') if usesandbox == True else os.getenv('PAYFAST_MERCHANT_ID')
    merchantkey = os.getenv('PAYFAST_SANDBOX_MERCHANT_KEY') if usesandbox == True else os.getenv('PAYFAST_MERCHANT_KEY')
    merchantpassphrase = os.getenv('PAYFAST_SANDBOX_MERCHANT_PASSPHRASE') if usesandbox == True else os.getenv('PAYFAST_MERCHANT_PASSPHRASE')  
    activationscript = str(os.getenv('PAYFAST_SANDBOX_MERCHANT_ACTIVESCRIPT')) if usesandbox == True else str(os.getenv('PAYFAST_MERCHANT_ACTIVESCRIPT')),  
    activationscriptstring = ''.join(activationscript)
    paymmethods = str(os.getenv('PAYFAST_SANDBOX_MERCHANT_PAYMMETHODS')) if usesandbox == True else str(os.getenv('PAYFAST_MERCHANT_PAYMMETHODS')),  

    routes = {
        'merchantid' : merchantid ,
        'merchantkey' : merchantkey ,
        'merchantpassphrase' : merchantpassphrase,     
        'activationscript' : activationscriptstring,
        'usesandbox' : usesandbox ,
        'paymmethods' : paymmethods
    }

    return Response(routes)


class EventsDetailsView(APIView):
    queryset = EventDetailTable.objects.all()
    serializer_class = EventDetailSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def get(self, request):
        body, error_response = _parse_body_or_400(request)
        if error_response:
            return error_response

        event_id = body.get('event') or request.data.get('event')
        if not event_id:
            return Response({'detail': 'event is required.'}, status=status.HTTP_400_BAD_REQUEST)

        events = EventDetailTable.objects.prefetch_related("EventSubDetailTable_set").filter(id=event_id)
        if not events.exists():
            logger.warning("EventsDetailsView event not found", extra={"event_id": event_id})
            return Response({'detail': 'Event not found.'}, status=status.HTTP_404_NOT_FOUND)

        serializer = EventDetailSerializer(events, many=True)
        return Response(serializer.data)
    
class NotifyUserView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        title = request.data.get("title", "Hello!")
        message = request.data.get("message", "You have a new alert.")
        event_id = request.data.get("event_id")
        notification_id = request.data.get("notification_id")

        if event_id is None or notification_id is None:
            logger.warning("NotifyUserView missing event_id or notification_id")
            return Response(
                {'detail': 'event_id and notification_id are required.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            result = send_push_notification(title, None, message, event_id, notification_id)
        except Exception:
            logger.exception(
                "NotifyUserView push failed",
                extra={"event_id": event_id, "notification_id": notification_id},
            )
            return Response(
                {'detail': 'Push notification failed.'},
                status=status.HTTP_502_BAD_GATEWAY,
            )

        if isinstance(result, dict) and result.get('error'):
            logger.error(
                "NotifyUserView push returned error",
                extra={
                    "event_id": event_id,
                    "notification_id": notification_id,
                    "error": result.get('error'),
                },
            )
            return Response(result, status=status.HTTP_502_BAD_GATEWAY)

        logger.info(
            "NotifyUserView push sent",
            extra={"event_id": event_id, "notification_id": notification_id},
        )
        return Response(result)
    

class EventNotificationViewSet(ModelViewSet):
    queryset = EventNotificationTable.objects.all()
    serializer_class = EventNotificationSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventNotificationTable.objects.all(),
            EventNotificationSerializer,
            ordering='-created',
        )
    
    def list(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventNotificationTable.objects.all(),
            EventNotificationSerializer,
            ordering='-created',
        )


def _organisation_notifications_from_body(body):
    queryset = OrganisationNotificationTable.objects.select_related(
        'athleticorganisation',
        'event',
    ).order_by('-created')

    if 'id' in body:
        return queryset.filter(id=body['id']), False
    if 'athleticorganisation' in body:
        return queryset.filter(athleticorganisation_id=body['athleticorganisation']), True
    return queryset, True


class OrganisationNotificationViewSet(ModelViewSet):
    queryset = OrganisationNotificationTable.objects.select_related(
        'athleticorganisation',
        'event',
    ).order_by('-created')
    serializer_class = OrganisationNotificationSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def retrieve(self, request, *args, **kwargs):
        body, error_response = _parse_body_or_400(request)
        if error_response:
            return error_response
        notifications, many = _organisation_notifications_from_body(body)
        serializer = OrganisationNotificationSerializer(notifications, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        body, error_response = _parse_body_or_400(request)
        if error_response:
            return error_response
        notifications, many = _organisation_notifications_from_body(body)
        serializer = OrganisationNotificationSerializer(notifications, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)
    

class EventSponsorViewSet(ModelViewSet):
    queryset = EventSponsorTable.objects.all()
    serializer_class = EventSponsorSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventSponsorTable.objects.all(),
            EventSponsorSerializer,
        )
    
    def list(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventSponsorTable.objects.all(),
            EventSponsorSerializer,
        )
    
class eventInfoViewSet(ModelViewSet):
    queryset = EventInformationTable.objects.all()
    serializer_class = EventInformationSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventInformationTable.objects.all(),
            EventInformationSerializer,
        )
    
    def list(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            EventInformationTable.objects.all(),
            EventInformationSerializer,
        )
        
class AppSponsorViewSet(ModelViewSet):
    queryset = AppSponsorTable.objects.all()
    serializer_class = AppSponsorSerializer
    parser_classes = (MultiPartParser , FormParser , JSONParser)

    def retrieve(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            AppSponsorTable.objects.all(),
            AppSponsorSerializer,
            event_field=None,
        )
    
    def list(self , request  , *args, **kwargs):
        return _filtered_viewset_response(
            request,
            AppSponsorTable.objects.all(),
            AppSponsorSerializer,
            event_field=None,
        )


class OrganisationViewSet(ModelViewSet):
    queryset = OrganisationTable.objects.all()
    serializer_class = OrganisationSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_organisations(self, request):
        body = _parse_request_body(request)
        queryset = OrganisationTable.objects.prefetch_related(
            'organisationmember_organisation', 'organisationevent_organisation'
        ).all()
        return _filter_queryset(queryset, body, {'id': 'id', 'status': 'status'})

    def retrieve(self, request, *args, **kwargs):
        organisations = self._get_organisations(request)
        many = 'id' not in _parse_request_body(request)
        serializer = OrganisationSerializer(organisations, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        organisations = self._get_organisations(request)
        serializer = OrganisationSerializer(organisations, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class OrganisationMemberViewSet(ModelViewSet):
    queryset = OrganisationMemberTable.objects.all()
    serializer_class = OrganisationMemberSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_organisation_members(self, request):
        body = _parse_request_body(request)
        queryset = OrganisationMemberTable.objects.select_related('organisation', 'participant').all()
        return _filter_queryset(
            queryset, body,
            {'id': 'id', 'organisation': 'organisation', 'participant': 'participant', 'status': 'status'},
        )

    def retrieve(self, request, *args, **kwargs):
        members = self._get_organisation_members(request)
        many = 'id' not in _parse_request_body(request)
        serializer = OrganisationMemberSerializer(members, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        members = self._get_organisation_members(request)
        serializer = OrganisationMemberSerializer(members, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class OrganisationEventViewSet(ModelViewSet):
    queryset = OrganisationEventTable.objects.all()
    serializer_class = OrganisationEventSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_organisation_events(self, request):
        body = _parse_request_body(request)
        queryset = OrganisationEventTable.objects.select_related('organisation', 'event').all()
        return _filter_queryset(
            queryset, body,
            {'id': 'id', 'organisation': 'organisation', 'event': 'event'},
        )

    def retrieve(self, request, *args, **kwargs):
        organisation_events = self._get_organisation_events(request)
        many = 'id' not in _parse_request_body(request)
        serializer = OrganisationEventSerializer(organisation_events, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        organisation_events = self._get_organisation_events(request)
        serializer = OrganisationEventSerializer(organisation_events, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class AthleticOrganisationViewSet(ModelViewSet):
    queryset = AthleticOrganisationTable.objects.all()
    serializer_class = AthleticOrganisationSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_athletic_organisations(self, request):
        body = _parse_request_body(request)
        queryset = AthleticOrganisationTable.objects.all()
        return _filter_queryset(
            queryset, body,
            {'id': 'id', 'code': 'code', 'organisationtype': 'organisationtype', 'isactive': 'isactive'},
        )

    def retrieve(self, request, *args, **kwargs):
        organisations = self._get_athletic_organisations(request)
        many = 'id' not in _parse_request_body(request)
        serializer = AthleticOrganisationSerializer(organisations, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        organisations = AthleticOrganisationTable.get_tree()
        serializer = AthleticOrganisationSerializer(organisations, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def create(self, request, *args, **kwargs):
        return Response(
            {
                'detail': 'Use /athleticorganisations/add_root/ or /athleticorganisations/add_child/.',
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

    def destroy(self, request, *args, **kwargs):
        return Response(
            {
                'detail': 'Use /athleticorganisations/delete/ with organisation_id in the request body.',
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )


@api_view(['GET'])
def athleticOrganisationTree(request):
    organisations = AthleticOrganisationTable.get_tree()
    serializer = AthleticOrganisationSerializer(organisations, many=True)
    return Response(serializer.data, status=status.HTTP_200_OK)


@api_view(['POST'])
def athleticOrganisationAddRoot(request):
    body = _parse_request_body(request)
    kwargs = _athletic_organisation_kwargs(body)
    missing = _validate_athletic_organisation_required_fields(kwargs)
    if missing:
        return Response(
            {'error': f'Required fields missing: {", ".join(missing)}'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if AthleticOrganisationTable.objects.filter(code=kwargs['code']).exists():
        logger.info("Athletic organisation add root rejected: duplicate code", extra={"code": kwargs['code']})
        return Response({'error': 'An organisation with this code already exists.'}, status=status.HTTP_409_CONFLICT)

    organisation = AthleticOrganisationTable.add_root(**kwargs)
    logger.info("Athletic organisation root created", extra={"organisation_id": organisation.id, "code": organisation.code})
    return Response(AthleticOrganisationSerializer(organisation).data, status=status.HTTP_201_CREATED)


@api_view(['POST'])
def athleticOrganisationAddChild(request):
    body = _parse_request_body(request)
    parent_id = body.get('parent_id')
    if not parent_id:
        return Response({'error': 'parent_id is required.'}, status=status.HTTP_400_BAD_REQUEST)

    parent = AthleticOrganisationTable.objects.filter(id=parent_id).first()
    if parent is None:
        return Response({'error': 'Parent organisation not found.'}, status=status.HTTP_404_NOT_FOUND)

    kwargs = _athletic_organisation_kwargs(body)
    missing = _validate_athletic_organisation_required_fields(kwargs)
    if missing:
        return Response(
            {'error': f'Required fields missing: {", ".join(missing)}'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if AthleticOrganisationTable.objects.filter(code=kwargs['code']).exists():
        logger.info("Athletic organisation add child rejected: duplicate code", extra={"code": kwargs['code']})
        return Response({'error': 'An organisation with this code already exists.'}, status=status.HTTP_409_CONFLICT)

    organisation = parent.add_child(**kwargs)
    logger.info(
        "Athletic organisation child created",
        extra={"organisation_id": organisation.id, "parent_id": parent.id, "code": organisation.code},
    )
    return Response(AthleticOrganisationSerializer(organisation).data, status=status.HTTP_201_CREATED)


@api_view(['POST'])
def athleticOrganisationMove(request):
    body = _parse_request_body(request)
    organisation_id = body.get('organisation_id')
    if not organisation_id:
        return Response({'error': 'organisation_id is required.'}, status=status.HTTP_400_BAD_REQUEST)

    organisation = AthleticOrganisationTable.objects.filter(id=organisation_id).first()
    if organisation is None:
        return Response({'error': 'Organisation not found.'}, status=status.HTTP_404_NOT_FOUND)

    target_parent_id = body.get('target_parent_id')
    position = body.get('position', 'sorted-child')
    target_parent = None
    if target_parent_id not in (None, ''):
        target_parent = AthleticOrganisationTable.objects.filter(id=target_parent_id).first()
        if target_parent is None:
            return Response({'error': 'Target parent organisation not found.'}, status=status.HTTP_404_NOT_FOUND)

    try:
        organisation = _move_athletic_organisation(organisation, target_parent, position)
    except ValueError as exc:
        logger.warning(
            "Athletic organisation move rejected",
            extra={"organisation_id": organisation_id, "error": str(exc)},
        )
        return Response({'error': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    logger.info(
        "Athletic organisation moved",
        extra={"organisation_id": organisation.id, "target_parent_id": target_parent_id, "position": position},
    )
    return Response(AthleticOrganisationSerializer(organisation).data, status=status.HTTP_200_OK)


@api_view(['POST'])
def athleticOrganisationUpdateNode(request):
    body = _parse_request_body(request)
    organisation_id = body.get('organisation_id') or body.get('id')
    if not organisation_id:
        return Response({'error': 'organisation_id is required.'}, status=status.HTTP_400_BAD_REQUEST)

    organisation = AthleticOrganisationTable.objects.filter(id=organisation_id).first()
    if organisation is None:
        return Response({'error': 'Organisation not found.'}, status=status.HTTP_404_NOT_FOUND)

    if 'code' in body and body['code'] and body['code'] != organisation.code:
        if AthleticOrganisationTable.objects.filter(code=body['code']).exclude(id=organisation.id).exists():
            return Response({'error': 'An organisation with this code already exists.'}, status=status.HTTP_409_CONFLICT)
        organisation.code = body['code']

    organisation = _update_athletic_organisation_fields(organisation, body)
    logger.info("Athletic organisation updated", extra={"organisation_id": organisation.id})
    return Response(AthleticOrganisationSerializer(organisation).data, status=status.HTTP_200_OK)


@api_view(['POST'])
def athleticOrganisationDeleteNode(request):
    body = _parse_request_body(request)
    organisation_id = body.get('organisation_id') or body.get('id')
    if not organisation_id:
        return Response({'error': 'organisation_id is required.'}, status=status.HTTP_400_BAD_REQUEST)

    organisation = AthleticOrganisationTable.objects.filter(id=organisation_id).first()
    if organisation is None:
        return Response({'error': 'Organisation not found.'}, status=status.HTTP_404_NOT_FOUND)

    organisation.delete()
    logger.warning("Athletic organisation deleted", extra={"organisation_id": organisation_id})
    return Response({'detail': 'Organisation deleted.'}, status=status.HTTP_200_OK)


class AthleticOrganisationMemberViewSet(ModelViewSet):
    queryset = AthleticOrganisationMemberTable.objects.all()
    serializer_class = AthleticOrganisationMemberSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_athletic_organisation_members(self, request):
        body = _parse_request_body(request)
        queryset = AthleticOrganisationMemberTable.objects.select_related(
            'athleticorganisation', 'participant',
        ).all()
        return _filter_queryset(
            queryset, body,
            {
                'id': 'id',
                'athleticorganisation': 'athleticorganisation',
                'participant': 'participant',
                'status': 'status',
            },
        )

    def retrieve(self, request, *args, **kwargs):
        members = self._get_athletic_organisation_members(request)
        many = 'id' not in _parse_request_body(request)
        serializer = AthleticOrganisationMemberSerializer(members, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        members = self._get_athletic_organisation_members(request)
        serializer = AthleticOrganisationMemberSerializer(members, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class AthleticOrganisationCategoryViewSet(ModelViewSet):
    queryset = AthleticOrganisationCategoryTable.objects.all()
    serializer_class = AthleticOrganisationCategorySerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    def _get_athletic_organisation_categories(self, request):
        body = _parse_request_body(request)
        queryset = AthleticOrganisationCategoryTable.objects.select_related(
            'athleticorganisation', 'athleticscategory',
        ).all()
        return _filter_queryset(
            queryset, body,
            {
                'id': 'id',
                'athleticorganisation': 'athleticorganisation',
                'athleticscategory': 'athleticscategory',
                'isactive': 'isactive',
            },
        )

    def retrieve(self, request, *args, **kwargs):
        categories = self._get_athletic_organisation_categories(request)
        many = 'id' not in _parse_request_body(request)
        serializer = AthleticOrganisationCategorySerializer(categories, many=many)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def list(self, request, *args, **kwargs):
        categories = self._get_athletic_organisation_categories(request)
        serializer = AthleticOrganisationCategorySerializer(categories, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class AppReleaseVersionViewSet(ModelViewSet):
    queryset = AppReleaseVersionTable.objects.all()
    serializer_class = AppReleaseVersionSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    @action(detail=False, methods=['post'], url_path='lookup')
    def lookup(self, request):
        body, error_response = _parse_body_or_400(request)
        if error_response:
            return error_response

        app_version = body.get('app_version')
        if not app_version:
            return Response(
                {'detail': 'app_version is required in request body.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        release = AppReleaseVersionTable.objects.filter(
            releaseVersionNumber=app_version,
        ).first()
        if release:
            logger.debug("App release version lookup hit", extra={"app_version": app_version})
            return Response(
                AppReleaseVersionSerializer(release).data,
                status=status.HTTP_200_OK,
            )
        logger.debug("App release version lookup miss", extra={"app_version": app_version})
        return Response(status=status.HTTP_204_NO_CONTENT)


class TermsAndConditionsViewSet(ModelViewSet):
    queryset = TermsAndConditionsTable.objects.all()
    serializer_class = TermsAndConditionsSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)

    @action(detail=False, methods=['get'], url_path='current')
    def current(self, request):
        document_type = request.query_params.get('document_type', 'app')
        current_terms = _get_current_terms(document_type)
        if not current_terms:
            return Response(
                {'detail': f'No current terms found for document_type "{document_type}".'},
                status=status.HTTP_204_NO_CONTENT,
            )
        return Response(TermsAndConditionsSerializer(current_terms).data)


@api_view(['GET'])
def getCurrentTerms(request):
    document_type = request.GET.get('document_type', 'app')
    current_terms = _get_current_terms(document_type)
    if not current_terms:
        return Response(
            {'detail': f'No current terms found for document_type "{document_type}".'},
            status=status.HTTP_204_NO_CONTENT,
        )
    return Response(TermsAndConditionsSerializer(current_terms).data)


@api_view(['POST'])
def getParticipantTermsStatus(request):
    participant, error_response = _get_participant_from_body(request)
    if error_response:
        return error_response

    body = _parse_request_body(request)
    document_type = body.get('document_type', 'app')
    logger.debug(
        "Participant terms status requested",
        extra={"participant_id": participant.id, "document_type": document_type},
    )
    status_data = _build_participant_terms_status(participant, document_type)
    return Response(ParticipantTermsStatusSerializer(status_data).data)


@api_view(['POST'])
def acceptParticipantTerms(request):
    participant, error_response = _get_participant_from_body(request)
    if error_response:
        return error_response

    body = _parse_request_body(request)
    document_type = body.get('document_type', 'app')
    terms_id = body.get('terms_id')

    if terms_id:
        terms = TermsAndConditionsTable.objects.filter(id=terms_id).first()
        if not terms:
            return Response(
                {'detail': 'Terms and conditions record not found.'},
                status=status.HTTP_404_NOT_FOUND,
            )
    else:
        terms = _get_current_terms(document_type)
        if not terms:
            return Response(
                {'detail': f'No current terms found for document_type "{document_type}".'},
                status=status.HTTP_404_NOT_FOUND,
            )

    if not terms.is_current:
        return Response(
            {'detail': 'Only the current terms and conditions version can be accepted.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if terms.document_type != document_type:
        return Response(
            {'detail': 'Terms document_type does not match request document_type.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    acceptance, created = ParticipantTermsAcceptanceTable.objects.get_or_create(
        participant=participant,
        terms=terms,
        defaults={
            'app_version': body.get('app_version'),
            'device_type': body.get('device_type'),
        },
    )
    if not created:
        logger.debug(
            "Participant terms already accepted",
            extra={"participant_id": participant.id, "terms_id": terms.id},
        )
        return Response(
            ParticipantTermsAcceptanceSerializer(acceptance).data,
            status=status.HTTP_200_OK,
        )
    logger.info(
        "Participant terms accepted",
        extra={"participant_id": participant.id, "terms_id": terms.id, "document_type": document_type},
    )
    return Response(
        ParticipantTermsAcceptanceSerializer(acceptance).data,
        status=status.HTTP_201_CREATED,
    )


class ErrorViewSet(ModelViewSet):
    queryset = ErrorTable.objects.all()
    serializer_class = ErrorSerializer
    parser_classes = (MultiPartParser, FormParser, JSONParser)
    http_method_names = ['get', 'post', 'head', 'options']

    def create(self, request, *args, **kwargs):
        try:
            response = super().create(request, *args, **kwargs)
        except Exception:
            logger.exception("Failed to save client error report")
            raise
        logger.info(
            "Client error report received",
            extra={"error_report_id": response.data.get('id')},
        )
        return response


@api_view(['GET'])
def users(request):
    return Response([
        {"id": 1, "name": "Deepak"},
        {"id": 2, "name": "John"}
    ])

class ParticipantImportView(APIView):

    def post(self, request):

        excel_file = request.FILES.get("file")

        if excel_file is None:
            logger.warning("Participant import rejected: no file uploaded")
            return Response(
                {"message":"No file selected"},
                status=400
            )

        logger.info(
            "Participant import started",
            extra={"filename": excel_file.name, "size_bytes": excel_file.size},
        )

        try:
            workbook = load_workbook(excel_file)
            worksheet = workbook["Entrants"]
        except KeyError:
            logger.error(
                "Participant import failed: worksheet 'Entrants' not found",
                extra={"filename": excel_file.name},
            )
            return Response(
                {"message": "Worksheet 'Entrants' not found in workbook."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except Exception:
            logger.exception(
                "Participant import failed while reading workbook",
                extra={"filename": excel_file.name},
            )
            return Response(
                {"message": "Could not read Excel file."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        imported = 0
        importedParticipant = 0
        updatedParticipant = 0
        errors = []

        try:
            usr = User.objects.get(username='Participant')
        except User.DoesNotExist:
            logger.error("Participant import failed: import user 'Participant' not found")
            return Response(
                {"message": "Import user 'Participant' is not configured."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        #
        # Skip header row
        #
        for row_number, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True),
                start=2):

            try:
                
                with transaction.atomic():
                    participant_id      = row[0]
                    event_id            = row[1]
                    event_name          = row[2]
                    first_name          = row[3]
                    last_name           = row[4]
                    category            = row[5]
                    distance_entered    = row[6]   
                    province            = row[7]
                    comp_shirt          = row[8]
                    email               = row[9]
                    cell_phone          = row[10]

                    #event = EventDetailTable.objects.get(id=event_id)
                    event = EventDetailTable.objects.filter(eventname=event_name).first()
                    category = EventCategoryTable.objects.filter(categoryname=category).first()

                    

                    participant, created = ParticipantTable.objects.update_or_create(
                        emailaddress=email,
                        defaults={
                            "firstname": first_name,
                            "surname": last_name,
                            "emailaddress": email,
                            "usrphonenum": cell_phone,
                            "user" : usr
                        }
                    )

                    subevent = EventSubDetailTable.objects.filter(event=event).filter(eventcategory=category).filter(name=distance_entered).first()

                    m_paym_id = f"imported-{participant.id}-{event.id}-{subevent.id}"[:30]
                    item_name = f"Imported from Excel - {event.eventname} - {subevent.name}"[:50]
                    payment_uuid = f"imported-{participant.id}-{event.id}-{subevent.id}"[:50]

                    current_datetime = timezone.now()
                    datetime_string = current_datetime.strftime("%Y-%m-%d %H:%M:%S")[:30]

                    participantPaymRefTable = ParticipantPaymRefTable.objects.create(
                        participant = participant,
                        m_paym_id = m_paym_id,
                        amount = 0.00,
                        item_name = item_name,
                        payment_uuid = payment_uuid,
                        payment_timestamp = datetime_string,
                        payment_status = 'Completed'
                    )

                    participantEvent = ParticipantEventTable.objects.create(
                        participant = participant,
                        event = event,
                        subevent = subevent,
                        paymref = participantPaymRefTable,
                        participantstatus = "Imported from Excel",
                        eventregdate = timezone.now()
                    )

                    if participantEvent is not None:
                        imported += 1

                    if created :
                        importedParticipant += 1
                    else:
                        updatedParticipant += 1

            except Exception as ex:
                error_message = f"Row {row_number}: {str(ex)}"
                errors.append(error_message)
                logger.warning(
                    "Participant import row failed",
                    extra={"row_number": row_number, "error": str(ex)},
                )

        logger.info(
            "Participant import completed",
            extra={
                "filename": excel_file.name,
                "imported": imported,
                "participants_created": importedParticipant,
                "participants_updated": updatedParticipant,
                "error_count": len(errors),
            },
        )

        return Response({
            "imported": imported,
            "participants_created": importedParticipant,
            "participants_updated": updatedParticipant,
            "errors": errors
        })